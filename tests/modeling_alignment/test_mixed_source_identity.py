"""Regression checks for actual document/table boundaries and safe rule defaults."""
from types import SimpleNamespace

import pytest

from kg_mnp.modeling.control_plane.errors import ModelingControlError
from kg_mnp.modeling.control_plane.mappings import table_identity
from kg_mnp.modeling.control_plane.providers.builtin import RuleMappingProvider
from tests.modeling_alignment.test_mapping_table_boundaries import ingest, mappings


def propose(dataset, classes=1):
    fields = mappings(dataset)
    for field in fields:
        field['target_property_iri'] = 'urn:property'
    return RuleMappingProvider().propose(SimpleNamespace(context={
        'kg_ir_items': dataset['items'], 'evidence_records': dataset['evidence_records'],
        'field_mappings': fields, 'default_namespace': 'urn:diagnostic:',
        'baseline_elements': [{'element_kind': 'CLASS', 'iri': f'urn:Class{i}',
                               'element_id': f'urn:class:{i}'} for i in range(classes)],
    }))


def test_docx_tables_have_independent_headers_and_entities(prompt03_workspace, tmp_path):
    from docx import Document
    document = Document()
    for header in ['left_field', 'right_field']:
        table = document.add_table(rows=2, cols=1)
        table.cell(0, 0).text = header
        table.cell(1, 0).text = 'value'
    path = tmp_path / 'tables.docx'
    document.save(path)
    dataset, _ = ingest(prompt03_workspace, [path])
    assert {row['source_field_name'] for row in mappings(dataset)} == {'left_field', 'right_field'}
    individuals = [d['body']['subject_iri'] for d in propose(dataset) if d['body']['candidate_type'] == 'INDIVIDUAL']
    assert len(individuals) == len(set(individuals)) == 2


def test_cross_source_row_numbers_never_define_identity(prompt03_workspace, tmp_path):
    files = [tmp_path / 'left.csv', tmp_path / 'right.csv']
    for index, path in enumerate(files):
        path.write_text(f'name,code\nT00{index},X{index}\n', encoding='utf-8')
    dataset, _ = ingest(prompt03_workspace, files)
    individuals = [d['body']['subject_iri'] for d in propose(dataset) if d['body']['candidate_type'] == 'INDIVIDUAL']
    assert len(individuals) == len(set(individuals)) == 2
    assert propose(dataset) == propose(dataset)
    assert not [d for d in propose(dataset, classes=2) if d['body']['candidate_type'] == 'CLASS_ASSERTION']


def test_table_coordinates_must_match_evidence(prompt03_workspace, tmp_path):
    path = tmp_path / 'rows.csv'
    path.write_text('name,code\nT001,X\n', encoding='utf-8')
    dataset, _ = ingest(prompt03_workspace, [path])
    cell = next(i for i in dataset['items'] if i['item_kind'] == 'table-cell')
    cell['payload']['row'] += 1
    with pytest.raises(ModelingControlError, match='coordinate'):
        table_identity(cell, {e['evidence_id']: e for e in dataset['evidence_records']})


def test_cross_sheet_rows_never_define_identity(prompt03_workspace, tmp_path):
    from openpyxl import Workbook
    book = Workbook()
    for sheet in [book.active, book.create_sheet('Other')]:
        sheet.append(['name'])
        sheet.append(['T001'])
    path = tmp_path / 'sheets.xlsx'
    book.save(path)
    dataset, _ = ingest(prompt03_workspace, [path])
    individuals = [d['body']['subject_iri'] for d in propose(dataset) if d['body']['candidate_type'] == 'INDIVIDUAL']
    assert len(individuals) == len(set(individuals)) == 2


def test_json_array_record_parents_are_not_the_whole_document(prompt03_workspace, tmp_path):
    path = tmp_path / 'records.json'
    path.write_text('[{"label":"Alpha","count":1},{"label":"Beta","count":2}]', encoding='utf-8')
    dataset, _ = ingest(prompt03_workspace, [path])
    drafts = propose(dataset)
    individuals = [d['body']['subject_iri'] for d in drafts if d['body']['candidate_type'] == 'INDIVIDUAL']
    assert len(individuals) == len(set(individuals)) == 2
    numbers = [d['body']['literal'] for d in drafts if d['body']['candidate_type'] == 'DATA_PROPERTY_ASSERTION' and d['body']['literal']['lexical_value'] in {'1', '2'}]
    assert len(numbers) == 2
    assert all(n['datatype_iri'].endswith('#integer') for n in numbers)

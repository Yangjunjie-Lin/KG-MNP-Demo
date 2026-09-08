"""Synthetic TXT + DOCX prose/tables + XLSX sheets through the real service."""
from __future__ import annotations

import asyncio
from copy import deepcopy
from io import BytesIO
from pathlib import Path

import pytest

from kg_mnp.jobs.worker import JobWorker
from kg_mnp.services.errors import ServiceBoundaryError
from kg_mnp.services.facade import ApplicationService
from kg_mnp.services.models import OperationRequest, ServiceConfiguration
from kg_mnp.services.uploads import receive_upload
from tests.services.test_modeling_workflow import call

MIN = 'https://yangjunjie-lin.github.io/KG-MNP-Demo/domain-packs/minimal/terms#'


def mixed_files():
    from docx import Document
    from openpyxl import Workbook
    document = Document()
    document.add_paragraph('实体 T002 的标签为 Beta。')
    table = document.add_table(rows=2, cols=2)
    for row, values in enumerate([['id', 'label'], ['legacy-A', 'Alpha']]):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    docx = BytesIO()
    document.save(docx)
    book = Workbook()
    book.active.title = 'Left'
    for sheet, values in [(book.active, ['T001', 'Alpha']), (book.create_sheet('Right'), ['T002', 'Beta'])]:
        sheet.append(['id', 'label'])
        sheet.append(values)
    xlsx = BytesIO()
    book.save(xlsx)
    return [('facts.txt', 'text/plain', '实体 T001 的标签为 Alpha。\n'.encode()),
            ('facts.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', docx.getvalue()),
            ('facts.xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', xlsx.getvalue())]


def profile(inventory):
    def common(identifier):
        return {'record_id': identifier, 'class_iri': MIN + 'Entity', 'identity_space': 'entities', 'id_field': 'id',
                'literals': {'label': {'predicate_iri': MIN + 'label', 'datatype': 'string'}}}
    rules = {'profile': 'evidence-record-mapping-v2', 'tables': [], 'text_templates': []}
    for index, table in enumerate(inventory['tables']):
        rules['tables'].append({**common(f'table-{index}'), 'source_id': table['source_id'], 'locator': table['locator'],
                               'identity_aliases': [{'value': 'legacy-A', 'canonical': 'T001', 'rationale': 'Synthetic cross-reference explicitly identifies the same test entity'}]})
    for index, source in enumerate(sorted({s for item in inventory['text_items'] for s in item['source_ids']})):
        rules['text_templates'].append({**common(f'text-{index}'), 'source_id': source, 'template': '实体 {id} 的标签为 {label}。'})
    return rules


@pytest.fixture
def mixed_case(tmp_path):
    root = Path(__file__).resolve().parents[2]
    service = ApplicationService(ServiceConfiguration(str(tmp_path / 'service'), review_profile='DEVELOPMENT_SINGLE_REVIEWER',
        reasoner_jar=str(root / 'third_party/downloads/robot-1.9.7.jar')))
    _, principal = service.tokens.create(principal_id='mixed-synthetic-reviewer', principal_type='HUMAN', permissions={'*'}, project_ids=set(), created_by='test')
    project_id = service.execute(OperationRequest('project.create', parameters={'name': 'mixed-source', 'domain_pack': 'minimal', 'domain_pack_version': '0.1.0'}), principal).payload['project_id']
    source_ids = []
    for filename, media, data in mixed_files():
        async def chunks(content=data):
            yield content
        upload = asyncio.run(receive_upload(service, project_id, principal, chunks(), filename=filename, media_type=media, idempotency_key=filename))
        job = JobWorker(service.jobs, service).run_once('mixed-upload')
        assert job.job_id == upload.job_id and job.status == 'SUCCEEDED', job.error
        source_ids.append(job.result['source']['source_id'])
    def execute(operation, params, key):
        return call(service, principal, project_id, operation, params, key)
    batch = execute('source.batch', {'source_ids': source_ids}, 'batch')['batch']
    plan = execute('ingestion.plan', {'batch_id': batch['batch_id']}, 'plan')['plan']
    run = execute('ingestion.run', {'plan_id': plan['plan_id']}, 'run')['run']
    scope = execute('modeling.scope', {'run_id': run['run_id'], 'description': 'Synthetic mixed-source identity and label assertions',
        'object_families': ['Entity'], 'in_scope': ['Entity labels'], 'namespace': 'urn:mixed:'}, 'scope')['scope']
    approval = execute('modeling.scope.approve', {'scope_id': scope['scope_id'], 'rationale': 'Bounded synthetic input'}, 'approve')['approval']
    prepared = execute('modeling.prepare', {'scope_id': scope['scope_id'], 'approval_id': approval['approval_id'],
        'questions': [{'question_text': 'Which entities and labels are present?', 'purpose': 'mixed-source deduplication', 'required_concepts': ['Entity']}]}, 'prepare')
    return service, principal, project_id, execute, prepared, source_ids


def propose(case, rules=None, key='propose'):
    _, _, _, execute, prepared, _ = case
    return execute('modeling.proposal', {'bundle_id': prepared['bundle']['modeling_input_bundle_id'],
        'providers': ['manual-candidate-provider'], 'record_mapping': rules or profile(prepared['input_inventory'])}, key)


def review(execute, proposed):
    head = None
    for index, item in enumerate(proposed['queue']['items']):
        if item['candidate_id']:
            action = execute('review.action', {'review_id': proposed['queue']['review_queue_id'], 'candidate_id': item['candidate_id'],
                'decision': 'ACCEPT', 'rationale': 'Explicit synthetic per-item human review of source identity and values', 'expected_head': head}, f'review-{index}')
            head = action['action']['action_hash']


def test_mixed_sources_deduplicate_with_all_evidence_and_real_compile(mixed_case):
    _, _, _, execute, prepared, sources = mixed_case
    proposed = propose(mixed_case)
    assert not proposed['extraction']['unmapped_item_ids']
    assert len(proposed['extraction']['text_spans']) == 4
    candidates = proposed['proposal']['abox_candidates']
    individuals = [c for c in candidates if c['body']['candidate_type'] == 'INDIVIDUAL']
    assert {c['body']['subject_iri'] for c in individuals} == {'urn:mixed:entities:T001', 'urn:mixed:entities:T002'}
    assert len(individuals) == 2
    assert not proposed['proposal']['conflicts']
    with pytest.raises(ServiceBoundaryError):
        execute('review.finalize', {'review_id': proposed['queue']['review_queue_id']}, 'premature-finalize')
    review(execute, proposed)
    confirmed = execute('review.finalize', {'review_id': proposed['queue']['review_queue_id']}, 'finalize')['confirmed_package']
    compilation = execute('compile.plan', {'confirmed_package_id': confirmed['package_id'], 'package_name': 'mixed-source', 'package_version': '0.1.0',
        'ontology_iri': 'urn:mixed:ontology', 'version_iri': 'urn:mixed:ontology:0.1.0',
        'oracles': [{'question_id': prepared['questions']['questions'][0]['question_id'], 'query_asset_id': 'minimal-query-list-entities',
                     'min_rows': 2, 'required_bindings': ['entity', 'label']}]}, 'compile-plan')['plan']
    built = execute('compile.build', {'plan_id': compilation['plan_id']}, 'build')
    assert built['reports']['shacl-validation-report.json']['status'] == 'CONFORMS'
    assert built['reports']['competency-question-test-report.json']['required_passed'] is True
    rows = execute('ods.query', {'package_id': built['package_id'], 'class_iri': MIN + 'Entity', 'limit': 100, 'offset': 0}, 'query')['rows']
    assert {row['iri'] for row in rows} == {'urn:mixed:entities:T001', 'urn:mixed:entities:T002'}
    traced = execute('object.trace', {'package_id': built['package_id'], 'instance_iri': 'urn:mixed:entities:T001'}, 'trace')
    assert {e['source_id'] for e in traced['evidence']} == set(sources)
    exported = execute('package.export', {'package_id': built['package_id']}, 'export')
    assert exported['size_bytes'] > 0


def test_unmapped_prose_blocks_confirmation_even_after_candidate_review(mixed_case):
    rules = profile(mixed_case[4]['input_inventory'])
    rules['text_templates'] = []
    proposed = propose(mixed_case, rules)
    assert len(proposed['extraction']['unmapped_item_ids']) == 2
    review(mixed_case[3], proposed)
    with pytest.raises(ServiceBoundaryError, match='Unmapped input'):
        mixed_case[3]('review.finalize', {'review_id': proposed['queue']['review_queue_id']}, 'finalize')


def test_conflicting_cross_source_labels_are_not_silently_merged(mixed_case):
    rules = profile(mixed_case[4]['input_inventory'])
    for table in rules['tables']:
        table['identity_aliases'].append({'value': 'T002', 'canonical': 'T001', 'rationale': 'Deliberately wrong synthetic merge to test conflict blocking'})
    proposed = propose(mixed_case, rules)
    assert {'LITERAL_CONFLICT', 'EVIDENCE_CONTRADICTION'} <= {c['code'] for c in proposed['proposal']['conflicts']}
    with pytest.raises(ServiceBoundaryError):
        mixed_case[3]('review.finalize', {'review_id': proposed['queue']['review_queue_id']}, 'finalize')


def test_mapping_sidecar_tamper_is_bound_to_reviewed_candidates(mixed_case):
    from kg_mnp.contracts.document_io import atomic_write_json, read_document
    from kg_mnp.modeling.control_plane.service import ModelingWorkspaceService
    from kg_mnp.services.projects import get_project
    service, _, project_id, execute, _, _ = mixed_case
    proposed = propose(mixed_case)
    review(execute, proposed)
    modeling = ModelingWorkspaceService(get_project(service.root, project_id).root)
    path = modeling.proposal_directory(proposed['proposal']['proposal_id']) / 'record-mapping-proposal.json'
    changed = deepcopy(read_document(path))
    changed['mapping']['tables'][0]['identity_space'] = 'tampered'
    atomic_write_json(path, changed)
    with pytest.raises(ServiceBoundaryError, match='Mapping no longer matches'):
        execute('review.finalize', {'review_id': proposed['queue']['review_queue_id']}, 'finalize')


def test_mapping_configuration_binds_provider_request_identity(mixed_case):
    from kg_mnp.contracts.canonical import semantic_hash
    from kg_mnp.contracts.document_io import read_document
    from kg_mnp.modeling.control_plane.service import ModelingWorkspaceService
    from kg_mnp.services.projects import get_project
    service, _, project_id, _, prepared, _ = mixed_case
    first = propose(mixed_case)
    rules = profile(prepared['input_inventory'])
    for definition in [*rules['tables'], *rules['text_templates']]:
        definition['identity_space'] = 'different-space'
    second = propose(mixed_case, rules, key='second-proposal')
    modeling = ModelingWorkspaceService(get_project(service.root, project_id).root)
    requests = []
    for proposed in [first, second]:
        directory = modeling.proposal_directory(proposed['proposal']['proposal_id'])
        request = read_document(directory/'provider-requests.json')[0]
        snapshot = read_document(directory/'provider-snapshots.json')[0]
        mapping = read_document(directory/'record-mapping-proposal.json')['mapping']
        assert snapshot['configuration_semantic_sha256'] == semantic_hash({'record_mapping': mapping})
        assert request['provider_snapshot_id'] == snapshot['snapshot_id']
        requests.append(request['request_id'])
    assert requests[0] != requests[1]

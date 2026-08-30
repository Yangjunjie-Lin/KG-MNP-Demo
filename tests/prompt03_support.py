from __future__ import annotations

import os
import subprocess
import sys
import wave
from io import BytesIO
from pathlib import Path


def run_cli(*arguments: str, cwd: Path | None = None) -> subprocess.CompletedProcess[str]:
    environment = os.environ.copy()
    environment["PYTHONUTF8"] = "1"
    return subprocess.run(
        [sys.executable, "-m", "kg_mnp.root_cli", *arguments],
        cwd=cwd,
        env=environment,
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
        timeout=60,
    )


def wav_bytes(*, frames: int = 80, sample_rate: int = 8000) -> bytes:
    stream = BytesIO()
    with wave.open(stream, "wb") as output:
        output.setnchannels(1)
        output.setsampwidth(1)
        output.setframerate(sample_rate)
        output.writeframes(b"\x80" * frames)
    return stream.getvalue()


def png_bytes() -> bytes:
    from PIL import Image

    stream = BytesIO()
    Image.new("RGB", (2, 3), color=(20, 40, 60)).save(stream, format="PNG")
    return stream.getvalue()


def xlsx_bytes() -> bytes:
    import openpyxl

    stream = BytesIO()
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Evidence"
    sheet["A1"] = "name"
    sheet["B1"] = "value"
    sheet["A2"] = "alpha"
    sheet["B2"] = "=1+1"
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def docx_bytes() -> bytes:
    from docx import Document

    stream = BytesIO()
    document = Document()
    document.add_paragraph("Evidence paragraph")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Evidence cell"
    document.save(stream)
    return stream.getvalue()


def pdf_bytes(text: str = "Evidence page") -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    result = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, value in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(f"{index} 0 obj\n".encode("ascii"))
        result.extend(value)
        result.extend(b"\nendobj\n")
    xref = len(result)
    result.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    result.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    result.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode(
            "ascii"
        )
    )
    return bytes(result)

"""Bounded Plugin SDK v1 conformance checks."""

from __future__ import annotations

import base64
import tempfile
import wave
from dataclasses import asdict, is_dataclass
from io import BytesIO
from pathlib import Path

from .api import (
    MediaDetectorPlugin,
    NormalizerPlugin,
    ParserPlugin,
    QualityEvaluatorPlugin,
    SourceAdapterPlugin,
)
from .errors import PluginConformanceError
from .models import (
    ConformanceResult,
    MediaDetection,
    MediaDetectionRequest,
    NormalizedUnit,
    NormalizeRequest,
    ParsedUnit,
    ParseRequest,
    PluginDescriptor,
    QualityEvaluationRequest,
    QualityEvaluationResult,
    ResourceLimits,
    SourceReadRequest,
    SourceReadResult,
)
from .registry import PluginRegistry
from .security import assert_safe_plugin_output
from .snapshot import build_snapshot, verify_snapshot


def _parser_sample(plugin_id: str) -> ParseRequest:
    limits = ResourceLimits()
    if plugin_id == "json-parser":
        return ParseRequest(b'{"value":1.25}', "application/json", limits)
    if plugin_id == "delimited-text-parser":
        return ParseRequest(b"name,value\nalpha,1\n", "text/csv", limits)
    if plugin_id == "wav-metadata-parser":
        stream = BytesIO()
        with wave.open(stream, "wb") as output:
            output.setnchannels(1)
            output.setsampwidth(1)
            output.setframerate(8000)
            output.writeframes(b"\x80" * 80)
        return ParseRequest(stream.getvalue(), "audio/wav", limits)
    if plugin_id == "image-metadata-parser":
        content = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        return ParseRequest(content, "image/png", limits)
    if plugin_id == "xlsx-parser":
        import openpyxl

        stream = BytesIO()
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = "value"
        workbook.save(stream)
        workbook.close()
        return ParseRequest(
            stream.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            limits,
        )
    if plugin_id == "docx-parser":
        from docx import Document

        stream = BytesIO()
        document = Document()
        document.add_paragraph("value")
        document.save(stream)
        return ParseRequest(
            stream.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            limits,
        )
    if plugin_id == "pdf-parser":
        from pypdf import PdfWriter

        stream = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        writer.write(stream)
        return ParseRequest(stream.getvalue(), "application/pdf", limits)
    media_type = "text/markdown" if plugin_id == "markdown-parser" else "text/plain"
    return ParseRequest(b"value\n", media_type, limits)


def _check_runtime_contract(plugin: object, descriptor: PluginDescriptor) -> None:
    kinds = set(descriptor.manifest["plugin_kinds"])
    if "source-adapter" in kinds:
        if not isinstance(plugin, SourceAdapterPlugin):
            raise PluginConformanceError("source adapter does not implement SourceAdapterPlugin")
        with tempfile.TemporaryDirectory(prefix="kg-mnp-plugin-conformance-") as directory:
            path = Path(directory) / "sample.txt"
            path.write_bytes(b"value\n")
            request = SourceReadRequest(path, ResourceLimits())
            first = plugin.read(request)
            second = plugin.read(request)
        if not isinstance(first, SourceReadResult) or first != second:
            raise PluginConformanceError("source adapter response is invalid or non-repeatable")
        assert_safe_plugin_output(asdict(first))
    if "media-detector" in kinds:
        if not isinstance(plugin, MediaDetectorPlugin):
            raise PluginConformanceError("media detector does not implement MediaDetectorPlugin")
        request = MediaDetectionRequest(b"value\n", "sample.txt")
        first = plugin.detect(request)
        second = plugin.detect(request)
        if not isinstance(first, MediaDetection) or first != second:
            raise PluginConformanceError("media detector response is invalid or non-repeatable")
        assert_safe_plugin_output(asdict(first))
    if "parser" in kinds:
        if not isinstance(plugin, ParserPlugin):
            raise PluginConformanceError("parser does not implement ParserPlugin")
        request = _parser_sample(descriptor.plugin_id)
        first = plugin.parse(request)
        second = plugin.parse(request)
        if (
            not isinstance(first, tuple)
            or not all(isinstance(item, ParsedUnit) for item in first)
            or first != second
        ):
            raise PluginConformanceError("parser response is invalid or non-repeatable")
        assert_safe_plugin_output(tuple(asdict(item) for item in first))
    if "normalizer" in kinds:
        if not isinstance(plugin, NormalizerPlugin):
            raise PluginConformanceError("normalizer does not implement NormalizerPlugin")
        sample = ParsedUnit(
            unit_kind="text-block",
            value="Cafe\u0301\r\n",
            locator={"locator_kind": "text-line-range", "start_line": 1, "end_line": 1},
            media_type="text/plain",
            ordinal=0,
        )
        first = plugin.normalize(NormalizeRequest(sample))
        second = plugin.normalize(NormalizeRequest(sample))
        if not isinstance(first, NormalizedUnit) or first != second:
            raise PluginConformanceError("normalizer response is invalid or non-repeatable")
        assert_safe_plugin_output(asdict(first))
    if "quality-evaluator" in kinds:
        if not isinstance(plugin, QualityEvaluatorPlugin):
            raise PluginConformanceError(
                "quality evaluator does not implement QualityEvaluatorPlugin"
            )
        request = QualityEvaluationRequest(())
        first = plugin.evaluate(request)
        second = plugin.evaluate(request)
        if not isinstance(first, QualityEvaluationResult) or first != second:
            raise PluginConformanceError(
                "quality evaluator response is invalid or non-repeatable"
            )
        assert_safe_plugin_output(asdict(first))


def run_conformance(
    registry: PluginRegistry,
    plugin_id: str,
) -> ConformanceResult:
    descriptor: PluginDescriptor = registry.get(plugin_id)
    checks: list[tuple[str, str]] = []
    errors: list[str] = []
    try:
        snapshot = build_snapshot(descriptor)
        verify_snapshot(descriptor, snapshot)
        checks.append(("manifest-and-snapshot", "PASS"))
        if descriptor.manifest["determinism"] == "NONDETERMINISTIC":
            raise PluginConformanceError("strict conformance rejects NONDETERMINISTIC provider")
        checks.append(("determinism-declaration", "PASS"))
        plugin = registry.load(plugin_id)
        if is_dataclass(plugin):
            assert_safe_plugin_output(asdict(plugin))
        checks.append(("load-after-enable", "PASS"))
        _check_runtime_contract(plugin, descriptor)
        checks.append(("request-response-contract", "PASS"))
        checks.append(("repeatable-response", "PASS"))
        checks.append(("authority-boundary", "PASS"))
    except Exception as exc:  # noqa: BLE001 - third-party Plugin exceptions are conformance failures
        errors.append(str(exc))
    return ConformanceResult(
        plugin_id=plugin_id,
        status="PASS" if not errors else "FAIL",
        checks=tuple(checks),
        errors=tuple(errors),
    )

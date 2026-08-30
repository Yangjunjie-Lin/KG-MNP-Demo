"""Read-only XLSX cell parser with ZIP/macro/external-relationship defenses."""

from io import BytesIO

from kg_mnp.plugins.builtin import reject_office_active_content, validate_zip_container
from kg_mnp.plugins.errors import PluginError
from kg_mnp.plugins.models import ParsedUnit, ParseRequest


class XlsxParser:
    def parse(self, request: ParseRequest) -> tuple[ParsedUnit, ...]:
        archive = validate_zip_container(request.content, request.limits)
        try:
            reject_office_active_content(archive)
        finally:
            archive.close()
        try:
            import openpyxl
        except ImportError as exc:
            raise PluginError("MISSING_DEPENDENCY: openpyxl") from exc
        try:
            workbook = openpyxl.load_workbook(
                BytesIO(request.content), read_only=True, data_only=False, keep_links=False
            )
        except Exception as exc:
            raise PluginError(f"invalid XLSX: {exc}") from exc
        units: list[ParsedUnit] = []
        try:
            if len(workbook.sheetnames) > request.limits.max_document_blocks:
                raise PluginError("XLSX_SHEET_LIMIT_EXCEEDED")
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                for row_index, row in enumerate(sheet.iter_rows(), 1):
                    if row_index > request.limits.max_table_rows:
                        raise PluginError("TABLE_ROW_LIMIT_EXCEEDED")
                    if len(row) > request.limits.max_table_columns:
                        raise PluginError("TABLE_COLUMN_LIMIT_EXCEEDED")
                    for column_index, cell in enumerate(row, 1):
                        if cell.value is None:
                            continue
                        value = str(cell.value)
                        if len(value) > request.limits.max_cell_characters:
                            raise PluginError("CELL_CHARACTER_LIMIT_EXCEEDED")
                        flags = ("FORMULA_PRESERVED_AS_TEXT",) if cell.data_type == "f" else ()
                        units.append(ParsedUnit("table-cell", value, {"locator_kind": "spreadsheet-cell", "sheet": sheet_name, "row": row_index, "column": column_index}, request.media_type, len(units), (("row", row_index), ("column", column_index), ("sheet", sheet_name)), flags))
                        if len(units) > request.limits.max_json_items:
                            raise PluginError("XLSX_CELL_LIMIT_EXCEEDED")
        finally:
            workbook.close()
        return tuple(units)

"""DOCX paragraph/table-cell parser with relationship and archive defenses."""

from io import BytesIO

from kg_mnp.plugins.builtin import reject_office_active_content, validate_zip_container
from kg_mnp.plugins.errors import PluginError
from kg_mnp.plugins.models import ParsedUnit, ParseRequest


class DocxParser:
    def parse(self, request: ParseRequest) -> tuple[ParsedUnit, ...]:
        archive = validate_zip_container(request.content, request.limits)
        try:
            reject_office_active_content(archive)
        finally:
            archive.close()
        try:
            from docx import Document
        except ImportError as exc:
            raise PluginError("MISSING_DEPENDENCY: python-docx") from exc
        try:
            document = Document(BytesIO(request.content))
        except Exception as exc:
            raise PluginError(f"invalid DOCX: {exc}") from exc
        units: list[ParsedUnit] = []
        for index, paragraph in enumerate(document.paragraphs):
            if len(units) >= request.limits.max_document_blocks:
                raise PluginError("DOCUMENT_BLOCK_LIMIT_EXCEEDED")
            if paragraph.text:
                units.append(ParsedUnit("text-block", paragraph.text, {"locator_kind": "document-paragraph", "paragraph_index": index}, request.media_type, len(units), (("block_kind", "paragraph"),)))
        for table_index, table in enumerate(document.tables):
            for row_index, row in enumerate(table.rows, 1):
                if row_index > request.limits.max_table_rows:
                    raise PluginError("TABLE_ROW_LIMIT_EXCEEDED")
                for column_index, cell in enumerate(row.cells, 1):
                    if column_index > request.limits.max_table_columns:
                        raise PluginError("TABLE_COLUMN_LIMIT_EXCEEDED")
                    if len(cell.text) > request.limits.max_cell_characters:
                        raise PluginError("CELL_CHARACTER_LIMIT_EXCEEDED")
                    units.append(ParsedUnit("table-cell", cell.text, {"locator_kind": "document-table-cell", "table_index": table_index, "row": row_index, "column": column_index}, request.media_type, len(units), (("row", row_index), ("column", column_index), ("table_index", table_index))))
                    if len(units) > request.limits.max_document_blocks:
                        raise PluginError("DOCUMENT_BLOCK_LIMIT_EXCEEDED")
        return tuple(units)
